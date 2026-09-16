import os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from database import SessionLocal
from init_db import SlidePresentation


def extraire_texte_shape(shape, profondeur=0) -> list:
    """
    Extrait le texte d'une shape, quel que soit son type :
    texte simple, tableau, forme groupée (récursif), ou texte alt d'image.
    Retourne une liste de morceaux de texte trouvés.
    """
    morceaux = []

    # 1. Texte simple (titre, zone de texte, forme avec texte)
    if shape.has_text_frame:
        texte = shape.text_frame.text.strip()
        if texte:
            morceaux.append(texte)

    # 2. Tableaux — ignorés par has_text_frame, il faut les traiter à part
    if shape.has_table:
        table = shape.table
        lignes_table = []
        for row in table.rows:
            cellules = [cell.text.strip() for cell in row.cells]
            # On ne garde la ligne que si elle contient au moins une info
            if any(cellules):
                lignes_table.append(" | ".join(cellules))
        if lignes_table:
            morceaux.append("[TABLEAU]\n" + "\n".join(lignes_table))

    # 3. Graphiques — on récupère les catégories et le nom des séries de données
    if shape.has_chart:
        try:
            chart = shape.chart
            infos_chart = []
            if chart.has_title and chart.chart_title.text_frame.text.strip():
                infos_chart.append(f"Titre graphique: {chart.chart_title.text_frame.text.strip()}")
            categories = [str(c) for c in chart.plots[0].categories] if chart.plots else []
            series_noms = [s.name for s in chart.series if s.name]
            if categories:
                infos_chart.append(f"Catégories: {', '.join(categories)}")
            if series_noms:
                infos_chart.append(f"Séries: {', '.join(series_noms)}")
            if infos_chart:
                morceaux.append("[GRAPHIQUE] " + " | ".join(infos_chart))
        except Exception:
            pass  # certains formats de graphique ne s'extraient pas proprement, on ignore sans planter

    # 4. Texte alternatif des images — souvent une description utile pour les schémas techniques
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        alt_text = ""
        try:
            alt_text = shape._element._nvXxPr.cNvPr.get("descr", "") or ""
        except Exception:
            pass
        if alt_text.strip():
            morceaux.append(f"[IMAGE - description: {alt_text.strip()}]")

    # 5. Formes groupées — on descend récursivement dedans, sinon leur contenu est invisible
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for sub_shape in shape.shapes:
            morceaux.extend(extraire_texte_shape(sub_shape, profondeur + 1))

    return morceaux


def extraire_notes(slide) -> str:
    """Extrait les notes du présentateur si présentes."""
    if slide.has_notes_slide:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
        return notes_text
    return ""


def import_powerpoint(chemin_fichier):
    # Vérification que le fichier existe
    if not os.path.exists(chemin_fichier):
        print(f"Erreur : Le fichier '{chemin_fichier}' est introuvable.")
        return

    nom_fichier = os.path.basename(chemin_fichier)
    print(f"Lecture du fichier PowerPoint '{nom_fichier}'...")

    # Chargement du fichier PPTX
    prs = Presentation(chemin_fichier)
    db = SessionLocal()

    imported = 0
    skipped = 0

    print(f"Nombre total de slides trouvées : {len(prs.slides)}\n")

    for index, slide in enumerate(prs.slides):
        numero = index + 1
        titre = ""
        contenu_complet = []

        try:
            # 1. Extraction du titre (s'il y en a un)
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                titre = slide.shapes.title.text.strip()

            # 2. Extraction de TOUTES les shapes (texte, tableaux, graphiques, groupes, images)
            for shape in slide.shapes:
                morceaux = extraire_texte_shape(shape)
                for morceau in morceaux:
                    # On évite de dupliquer le titre déjà extrait à part
                    if morceau and morceau != titre:
                        contenu_complet.append(morceau)

            # 3. Extraction des notes du présentateur (souvent riches en détails)
            notes = extraire_notes(slide)
            if notes:
                contenu_complet.append(f"[NOTES DU PRÉSENTATEUR]\n{notes}")

            # On rassemble tout le texte de la slide avec des sauts de ligne
            texte_final = "\n\n".join(contenu_complet)

            # 4. Vérification des doublons (Même fichier, même numéro de slide)
            existing = db.query(SlidePresentation).filter(
                SlidePresentation.nom_fichier == nom_fichier,
                SlidePresentation.numero_slide == numero
            ).first()

            if existing:
                skipped += 1
                continue

            # 5. Création et sauvegarde dans la base de données
            nouvelle_slide = SlidePresentation(
                nom_fichier=nom_fichier,
                numero_slide=numero,
                titre_slide=titre,
                contenu_texte=texte_final
            )

            db.add(nouvelle_slide)
            db.commit()  # Sauvegarde immédiate
            imported += 1

            print(f"Slide {numero} importée (Titre: {titre[:30]}...) — {len(contenu_complet)} bloc(s) de contenu")

        except Exception as e:
            print(f"Erreur sur la slide {numero} : {e}")
            db.rollback()

    db.close()
    print(f"\n Import terminé ! Importées : {imported} | Ignorées (déjà existantes) : {skipped}")


if __name__ == "__main__":
    fichier_ppt = input("Chemin du fichier PowerPoint (.pptx) : ")
    import_powerpoint(fichier_ppt)